class DocumentLoader:
    """
    Loads non-tabular files (PDF, Word .docx, images) and extracts
    text content so it can be previewed and passed to the AI as context.
    """

    def __init__(self):
        self.kind = None       # "pdf" | "docx" | "image"
        self.filename = ""
        self.text = ""
        self.image = None
        self.page_count = 0
        self.ocr_available = True

    def load(self, uploaded_file):

        if uploaded_file is None:
            return None

        self.filename = uploaded_file.name
        filename = uploaded_file.name.lower()

        if filename.endswith(".pdf"):
            self._load_pdf(uploaded_file)
        elif filename.endswith(".docx"):
            self._load_docx(uploaded_file)
        elif filename.endswith((".png", ".jpg", ".jpeg")):
            self._load_image(uploaded_file)

        return self.text

    def _load_pdf(self, uploaded_file):
        from pypdf import PdfReader

        self.kind = "pdf"
        reader = PdfReader(uploaded_file)
        self.page_count = len(reader.pages)

        pages_text = []
        for page in reader.pages:
            pages_text.append(page.extract_text() or "")

        self.text = "\n".join(pages_text).strip()

    def _load_docx(self, uploaded_file):
        import docx

        self.kind = "docx"
        document = docx.Document(uploaded_file)

        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        # also pull text out of any tables in the document
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)

        self.text = "\n".join(paragraphs).strip()

    def _load_image(self, uploaded_file):
        from PIL import Image

        self.kind = "image"
        self.image = Image.open(uploaded_file)

        try:
            import pytesseract
            self.text = pytesseract.image_to_string(self.image).strip()
        except Exception:
            # Tesseract binary not installed on this machine, or OCR failed.
            self.ocr_available = False
            self.text = ""

    def get_word_count(self):
        return len(self.text.split()) if self.text else 0

    def get_preview(self, max_chars=1500):
        if not self.text:
            return "No extractable text found in this file."

        preview = self.text[:max_chars]
        if len(self.text) > max_chars:
            preview += "..."

        return preview

    def get_context(self, max_chars=6000):

        if not self.text:
            if self.kind == "image":
                return (
                    f"An image file named '{self.filename}' was uploaded. "
                    "No text could be extracted from it "
                    "(OCR may be unavailable, or the image has no text)."
                )
            return (
                f"A {self.kind} file named '{self.filename}' was uploaded, "
                "but no extractable text was found."
            )

        context = self.text[:max_chars]

        return f"""
Uploaded Document: {self.filename}
Type: {self.kind.upper() if self.kind else "UNKNOWN"}

Content:

{context}
"""
